"""
exporters.py
--------------
Converts a ParsedProgram into documentation files.

CSV and HTML export use only the standard library and always work.
Excel (.xlsx) export requires openpyxl; Word (.docx) export requires
python-docx. Both are optional - if not installed, the corresponding
export function raises ExportDependencyError with an actionable
message instead of crashing, so the GUI can show it to the user.
"""

import csv
import html
import os
import time


class ExportDependencyError(Exception):
    """Raised when an optional export dependency is missing."""
    pass


def _timestamp():
    return time.strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# CSV export (always available)
# ---------------------------------------------------------------------------

def export_csv(parsed, output_dir):
    """Writes three CSV files: tags, data_types (flattened members), routines."""
    os.makedirs(output_dir, exist_ok=True)
    paths = {}

    tags_path = os.path.join(output_dir, "tags.csv")
    with open(tags_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Name", "Tag Type", "Data Type", "External Access", "Description"])
        for tag in parsed.tags:
            writer.writerow([tag["name"], tag["tag_type"], tag["data_type"],
                              tag["external_access"], tag["description"]])
    paths["tags"] = tags_path

    dtypes_path = os.path.join(output_dir, "data_types.csv")
    with open(dtypes_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["DataType", "DataType Description", "Member",
                          "Member DataType", "Member Access", "Member Description"])
        for dt in parsed.data_types:
            if not dt["members"]:
                writer.writerow([dt["name"], dt["description"], "", "", "", ""])
            for member in dt["members"]:
                writer.writerow([dt["name"], dt["description"], member["name"],
                                  member["data_type"], member["external_access"],
                                  member["description"]])
    paths["data_types"] = dtypes_path

    routines_path = os.path.join(output_dir, "routines.csv")
    with open(routines_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Program", "Routine", "Type", "Description"])
        for r in parsed.routines:
            writer.writerow([r["program"], r["name"], r["type"], r["description"]])
    paths["routines"] = routines_path

    return paths


# ---------------------------------------------------------------------------
# HTML export (always available)
# ---------------------------------------------------------------------------

def export_html(parsed, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, "documentation.html")

    def esc(value):
        return html.escape(str(value)) if value else ""

    rows_tags = "".join(
        f"<tr><td>{esc(t['name'])}</td><td>{esc(t['tag_type'])}</td>"
        f"<td>{esc(t['data_type'])}</td><td>{esc(t['external_access'])}</td>"
        f"<td>{esc(t['description'])}</td></tr>"
        for t in parsed.tags
    )

    dtype_blocks = ""
    for dt in parsed.data_types:
        member_rows = "".join(
            f"<tr><td>{esc(m['name'])}</td><td>{esc(m['data_type'])}</td>"
            f"<td>{esc(m['external_access'])}</td><td>{esc(m['description'])}</td></tr>"
            for m in dt["members"]
        )
        dtype_blocks += f"""
        <h3>{esc(dt['name'])}</h3>
        <p class="desc">{esc(dt['description'])}</p>
        <table>
          <tr><th>Member</th><th>Data Type</th><th>Access</th><th>Description</th></tr>
          {member_rows if member_rows else '<tr><td colspan="4"><em>No members</em></td></tr>'}
        </table>
        """

    rows_routines = "".join(
        f"<tr><td>{esc(r['program'])}</td><td>{esc(r['name'])}</td>"
        f"<td>{esc(r['type'])}</td><td>{esc(r['description'])}</td></tr>"
        for r in parsed.routines
    )

    html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>PLC Documentation — {esc(parsed.controller_name)}</title>
<style>
  body {{ font-family: Segoe UI, Arial, sans-serif; margin: 40px; color: #1f2937; }}
  h1 {{ color: #1e3a8a; }}
  h2 {{ border-bottom: 2px solid #1e3a8a; padding-bottom: 4px; margin-top: 40px; }}
  h3 {{ color: #1e3a8a; margin-top: 24px; }}
  table {{ border-collapse: collapse; width: 100%; margin-top: 8px; }}
  th, td {{ border: 1px solid #d1d5db; padding: 6px 10px; text-align: left; font-size: 14px; }}
  th {{ background-color: #eff6ff; }}
  tr:nth-child(even) {{ background-color: #f9fafb; }}
  .meta {{ color: #6b7280; font-size: 13px; }}
  .desc {{ color: #374151; font-style: italic; margin: 4px 0 8px 0; }}
</style>
</head>
<body>
  <h1>PLC Documentation</h1>
  <p class="meta">Controller: <strong>{esc(parsed.controller_name)}</strong>
     &nbsp;|&nbsp; Software revision: {esc(parsed.software_revision)}
     &nbsp;|&nbsp; Generated: {_timestamp()}</p>

  <h2>Tags ({len(parsed.tags)})</h2>
  <table>
    <tr><th>Name</th><th>Tag Type</th><th>Data Type</th><th>External Access</th><th>Description</th></tr>
    {rows_tags if rows_tags else '<tr><td colspan="5"><em>No tags found</em></td></tr>'}
  </table>

  <h2>User-Defined Data Types ({len(parsed.data_types)})</h2>
  {dtype_blocks if dtype_blocks else '<p><em>No user-defined data types found</em></p>'}

  <h2>Routines ({len(parsed.routines)})</h2>
  <table>
    <tr><th>Program</th><th>Routine</th><th>Type</th><th>Description</th></tr>
    {rows_routines if rows_routines else '<tr><td colspan="4"><em>No routines found</em></td></tr>'}
  </table>
</body>
</html>
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(html_doc)
    return path


# ---------------------------------------------------------------------------
# Excel export (optional, requires openpyxl)
# ---------------------------------------------------------------------------

def export_excel(parsed, output_dir):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise ExportDependencyError(
            "Excel export requires the 'openpyxl' package.\n"
            "In Thonny: Tools > Manage packages > search 'openpyxl' > Install."
        )

    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, "documentation.xlsx")

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    def style_header(ws, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

    ws_tags = wb.active
    ws_tags.title = "Tags"
    ws_tags.append(["Name", "Tag Type", "Data Type", "External Access", "Description"])
    for tag in parsed.tags:
        ws_tags.append([tag["name"], tag["tag_type"], tag["data_type"],
                         tag["external_access"], tag["description"]])
    style_header(ws_tags, 5)

    ws_dtypes = wb.create_sheet("Data Types")
    ws_dtypes.append(["DataType", "DataType Description", "Member",
                       "Member DataType", "Member Access", "Member Description"])
    for dt in parsed.data_types:
        if not dt["members"]:
            ws_dtypes.append([dt["name"], dt["description"], "", "", "", ""])
        for member in dt["members"]:
            ws_dtypes.append([dt["name"], dt["description"], member["name"],
                               member["data_type"], member["external_access"],
                               member["description"]])
    style_header(ws_dtypes, 6)

    ws_routines = wb.create_sheet("Routines")
    ws_routines.append(["Program", "Routine", "Type", "Description"])
    for r in parsed.routines:
        ws_routines.append([r["program"], r["name"], r["type"], r["description"]])
    style_header(ws_routines, 4)

    for ws in wb.worksheets:
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 60)

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Word export (optional, requires python-docx)
# ---------------------------------------------------------------------------

def export_word(parsed, output_dir):
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise ExportDependencyError(
            "Word export requires the 'python-docx' package.\n"
            "In Thonny: Tools > Manage packages > search 'python-docx' > Install."
        )

    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, "documentation.docx")

    document = docx.Document()
    title = document.add_heading("PLC Documentation", level=0)

    meta = document.add_paragraph()
    meta.add_run(f"Controller: {parsed.controller_name}\n").bold = True
    meta.add_run(f"Software revision: {parsed.software_revision}\n")
    meta.add_run(f"Generated: {_timestamp()}")

    document.add_heading(f"Tags ({len(parsed.tags)})", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, text in enumerate(["Name", "Tag Type", "Data Type", "External Access", "Description"]):
        hdr[i].text = text
    for tag in parsed.tags:
        row = table.add_row().cells
        row[0].text = tag["name"]
        row[1].text = tag["tag_type"]
        row[2].text = tag["data_type"]
        row[3].text = tag["external_access"]
        row[4].text = tag["description"]

    document.add_heading(f"User-Defined Data Types ({len(parsed.data_types)})", level=1)
    for dt in parsed.data_types:
        document.add_heading(dt["name"], level=2)
        if dt["description"]:
            p = document.add_paragraph()
            p.add_run(dt["description"]).italic = True
        if dt["members"]:
            mtable = document.add_table(rows=1, cols=4)
            mtable.style = "Light Grid Accent 1"
            mhdr = mtable.rows[0].cells
            for i, text in enumerate(["Member", "Data Type", "Access", "Description"]):
                mhdr[i].text = text
            for member in dt["members"]:
                row = mtable.add_row().cells
                row[0].text = member["name"]
                row[1].text = member["data_type"]
                row[2].text = member["external_access"]
                row[3].text = member["description"]

    document.add_heading(f"Routines ({len(parsed.routines)})", level=1)
    rtable = document.add_table(rows=1, cols=4)
    rtable.style = "Light Grid Accent 1"
    rhdr = rtable.rows[0].cells
    for i, text in enumerate(["Program", "Routine", "Type", "Description"]):
        rhdr[i].text = text
    for r in parsed.routines:
        row = rtable.add_row().cells
        row[0].text = r["program"]
        row[1].text = r["name"]
        row[2].text = r["type"]
        row[3].text = r["description"]

    document.save(path)
    return path
